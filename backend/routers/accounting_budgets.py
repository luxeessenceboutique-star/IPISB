import csv
import io
import unicodedata
from datetime import date, datetime, timezone
from typing import Annotated, Optional

from fastapi import APIRouter, Depends, HTTPException, Response, UploadFile
from fastapi.responses import StreamingResponse
from supabase import Client

from deps import get_current_user, get_db, CurrentUser
from models import BudgetCreate, BudgetUpdate
from utils.audit import log_audit
from utils.excel import make_xlsx

router = APIRouter(prefix="/accounting/budgets", tags=["accounting"])

MONTHS_FR = ["Janvier", "Février", "Mars", "Avril", "Mai", "Juin", "Juillet",
             "Août", "Septembre", "Octobre", "Novembre", "Décembre"]


def _require_admin(user: CurrentUser) -> None:
    if not user.can_access_accounting_full():
        raise HTTPException(403, "Admin only")


def _shape(b: dict) -> dict:
    return {
        **{k: v for k, v in b.items() if k != "accounting_categories"},
        "category_name": (b.get("accounting_categories") or {}).get("name"),
    }


# ── Période : année entière / mois précis / plage « du … au … » ──────────────
def _norm(s) -> str:
    """minuscule, sans accent, espaces normalisés — pour comparer libellés."""
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode()
    return " ".join(s.lower().split())


_MONTH_IDX = {_norm(m): i + 1 for i, m in enumerate(MONTHS_FR)}


def _parse_date(v) -> Optional[date]:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y", "%Y/%m/%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def _parse_month(v) -> Optional[int]:
    if v is None or v == "":
        return None
    if isinstance(v, bool):
        raise ValueError("mois invalide")
    if isinstance(v, (int, float)):
        return int(v)
    s = str(v).strip()
    if s.isdigit():
        return int(s)
    n = _norm(s)
    if n in _MONTH_IDX:
        return _MONTH_IDX[n]
    if n in ("annee entiere", "annuel", "annee", "toute l annee", "an"):
        return None
    raise ValueError(f"mois non reconnu : « {s} »")


def _parse_amount(v) -> float:
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    s = str(v)
    for junk in ("MAD", "mad", "DH", "dh", "Dh"):
        s = s.replace(junk, "")
    s = "".join(s.split())             # retire tous les espaces (dont insecables)
    if "," in s and "." in s:          # « 1.000,50 » → « 1000.50 »
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:                     # « 1000,50 » → « 1000.50 »
        s = s.replace(",", ".")
    try:
        return round(float(s), 2)
    except ValueError:
        raise ValueError(f"montant invalide : « {v} »")


def _period_or_raise(month, start_date, end_date):
    """Retourne (month, start_iso|None, end_iso|None). Lève ValueError si incohérent."""
    has_range = bool(start_date) or bool(end_date)
    if has_range:
        d1, d2 = _parse_date(start_date), _parse_date(end_date)
        if not d1 or not d2:
            raise ValueError("plage de dates incomplète ou invalide (format attendu AAAA-MM-JJ)")
        if d1 > d2:
            raise ValueError("la date de début doit précéder la date de fin")
        return None, d1.isoformat(), d2.isoformat()
    if month is not None:
        month = int(month)
        if not (1 <= month <= 12):
            raise ValueError("le mois doit être compris entre 1 et 12")
    return month, None, None


def _clean_period(month, start_date, end_date):
    try:
        return _period_or_raise(month, start_date, end_date)
    except ValueError as ex:
        raise HTTPException(400, str(ex))


def _period_label(b: dict) -> str:
    if b.get("start_date"):
        return f"Du {_fr(b['start_date'])} au {_fr(b.get('end_date'))}"
    m = b.get("month")
    return "Année entière" if m is None else MONTHS_FR[m - 1]


def _fr(iso: Optional[str]) -> str:
    d = _parse_date(iso)
    return d.strftime("%d/%m/%Y") if d else "…"


def _write_error(ex: Exception) -> HTTPException:
    """Traduit une erreur PostgREST en réponse HTTP lisible."""
    msg = str(ex)
    if any(t in msg for t in ("start_date", "end_date", "budgets_date_range_chk",
                              "budgets_range_no_month_chk", "does not exist",
                              "Could not find", "PGRST204", "42703")):
        return HTTPException(400, "Migration L42 requise (colonnes start_date / end_date des budgets).")
    if "duplicate key" in msg or "23505" in msg or "unique" in msg.lower():
        return HTTPException(409, "Un budget existe déjà pour cette catégorie / période.")
    return HTTPException(500, "Écriture du budget impossible.")


@router.get("")
async def list_budgets(
    user: Annotated[CurrentUser, Depends(get_current_user)],
    db: Annotated[Client, Depends(get_db)],
    year: Optional[int] = None,
    category_id: Optional[str] = None,
):
    _require_admin(user)
    query = db.from_("budgets").select("*, accounting_categories(name)")
    if year is not None:
        query = query.eq("year", year)
    if category_id:
        query = query.eq("category_id", category_id)
    rows = query.order("year", desc=True).order("month").execute().data or []
    return [_shape(b) for b in rows]


# ── Export du tableau (Excel / CSV) ─────────────────────────────────────────
def _export_rows(db: Client, year: Optional[int]) -> list[dict]:
    q = db.from_("budgets").select("*, accounting_categories(name)")
    if year is not None:
        q = q.eq("year", year)
    raw = q.order("year", desc=True).order("month").execute().data or []
    out = []
    for b in raw:
        out.append({
            "category": (b.get("accounting_categories") or {}).get("name") or "",
            "year": b.get("year"),
            "month": MONTHS_FR[b["month"] - 1] if b.get("month") else "",
            "start_date": b.get("start_date") or "",
            "end_date": b.get("end_date") or "",
            "period": _period_label(b),
            "amount": float(b.get("amount") or 0),
            "comment": b.get("comment") or "",
        })
    return out


@router.get("/export")
async def export_budgets(
    user: Annotated[CurrentUser, Depends(get_current_user)],
    db: Annotated[Client, Depends(get_db)],
    year: Optional[int] = None,
    fmt: str = "xlsx",
):
    _require_admin(user)
    rows = _export_rows(db, year)
    today = datetime.now(timezone.utc).date()
    total = sum(r["amount"] for r in rows)
    scope = f"exercice {year}" if year else "tous exercices"
    stem = f"Budgets_{year or 'tous'}"

    if fmt == "csv":
        buf = io.StringIO()
        w = csv.writer(buf, delimiter=";")
        w.writerow(["Catégorie", "Année", "Mois", "Date début", "Date fin",
                    "Montant prévu", "Commentaire"])
        for r in rows:
            w.writerow([r["category"], r["year"], r["month"], r["start_date"],
                        r["end_date"], f'{r["amount"]:.2f}', r["comment"]])
        return StreamingResponse(
            io.BytesIO(buf.getvalue().encode("utf-8-sig")),
            media_type="text/csv",
            headers={"Content-Disposition": f'attachment; filename="{stem}.csv"'},
        )

    return make_xlsx(
        filename=f"{stem}.xlsx",
        title="BUDGET PRÉVISIONNEL IPISB",
        subtitle=f"Édité le {today.strftime('%d/%m/%Y')} — {scope} — "
                 f"{len(rows)} ligne(s) — Total : {total:,.2f} MAD".replace(",", " "),
        theme="green",
        sheet_name="Budgets",
        columns=[
            {"key": "category", "label": "Catégorie", "width": 28},
            {"key": "year", "label": "Année", "type": "int", "width": 10},
            {"key": "month", "label": "Mois", "width": 13},
            {"key": "start_date", "label": "Date début", "type": "date", "width": 13},
            {"key": "end_date", "label": "Date fin", "type": "date", "width": 13},
            {"key": "amount", "label": "Montant prévu", "type": "money", "width": 16},
            {"key": "comment", "label": "Commentaire", "width": 34},
        ],
        rows=rows,
    )


@router.get("/import/template")
async def import_template(
    user: Annotated[CurrentUser, Depends(get_current_user)],
    db: Annotated[Client, Depends(get_db)],
    fmt: str = "xlsx",
):
    """Modèle vierge à remplir puis réimporter. Les libellés de catégorie
    doivent correspondre exactement à ceux de la liste des catégories."""
    _require_admin(user)
    cats = db.from_("accounting_categories").select("name").order("name").execute().data or []
    cat_list = ", ".join(c["name"] for c in cats) or "—"
    headers = ["Catégorie", "Année", "Mois", "Date début", "Date fin",
               "Montant prévu", "Commentaire"]

    if fmt == "csv":
        buf = io.StringIO()
        w = csv.writer(buf, delimiter=";")
        w.writerow(headers)
        w.writerow(["# Catégories valides : " + cat_list])
        return StreamingResponse(
            io.BytesIO(buf.getvalue().encode("utf-8-sig")),
            media_type="text/csv",
            headers={"Content-Disposition": 'attachment; filename="Budgets_modele.csv"'},
        )

    return make_xlsx(
        filename="Budgets_modele.xlsx",
        title="MODÈLE D'IMPORT — BUDGET PRÉVISIONNEL",
        subtitle="Une ligne par budget. Mois OU plage de dates OU rien (= année entière). "
                 f"Catégories valides : {cat_list}",
        theme="green",
        sheet_name="Budgets",
        columns=[
            {"key": "category", "label": "Catégorie", "width": 28},
            {"key": "year", "label": "Année", "width": 10},
            {"key": "month", "label": "Mois", "width": 13},
            {"key": "start_date", "label": "Date début", "width": 13},
            {"key": "end_date", "label": "Date fin", "width": 13},
            {"key": "amount", "label": "Montant prévu", "width": 16},
            {"key": "comment", "label": "Commentaire", "width": 34},
        ],
        rows=[],
    )


# ── Import du tableau (Excel / CSV) ─────────────────────────────────────────
_COL_ALIASES = {
    "category": ("categorie", "cat"),
    "year": ("annee", "exercice"),
    "month": ("mois",),
    "start_date": ("date debut", "debut", "du", "date de debut"),
    "end_date": ("date fin", "fin", "au", "date de fin"),
    "amount": ("montant", "montant prevu", "budget", "montant mad"),
    "comment": ("commentaire", "note", "remarque"),
}


def _map_headers(cells) -> dict[int, str]:
    mapping: dict[int, str] = {}
    for idx, cell in enumerate(cells):
        hn = _norm(cell)
        if not hn:
            continue
        for key, aliases in _COL_ALIASES.items():
            cands = (key, *aliases)
            # égalité stricte, ou préfixe pour les libellés longs (≥ 4 car.)
            # afin d'éviter qu'un « Durée » ne soit pris pour « du ».
            if hn in cands or any(len(c) >= 4 and hn.startswith(c) for c in cands):
                mapping.setdefault(idx, key)
                break
    return mapping


def _rows_to_records(all_rows: list) -> list[dict]:
    header_i = None
    mapping: dict[int, str] = {}
    for i, row in enumerate(all_rows[:10]):
        if not row:
            continue
        m = _map_headers(row)
        if "category" in m.values() and len(m) >= 2:
            header_i, mapping = i, m
            break
    if header_i is None:
        raise HTTPException(400, "En-têtes introuvables — le fichier doit contenir une colonne « Catégorie ».")

    records = []
    for row in all_rows[header_i + 1:]:
        if not row or not any(c is not None and str(c).strip() for c in row):
            continue
        rec = {key: (row[idx] if idx < len(row) else None) for idx, key in mapping.items()}
        records.append(rec)
    return records


def _parse_xlsx(raw: bytes) -> list[dict]:
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    except Exception:
        raise HTTPException(400, "Fichier Excel illisible.")
    ws = wb.active
    all_rows = [list(r) for r in ws.iter_rows(values_only=True)]
    return _rows_to_records(all_rows)


def _parse_csv(raw: bytes) -> list[dict]:
    text = raw.decode("utf-8-sig", errors="replace")
    first = next((ln for ln in text.splitlines() if ln.strip()), "")
    delim = ";" if first.count(";") >= first.count(",") else ","
    all_rows = [r for r in csv.reader(io.StringIO(text), delimiter=delim)]
    # ignore les lignes de commentaire « # … »
    all_rows = [r for r in all_rows if not (r and str(r[0]).lstrip().startswith("#"))]
    return _rows_to_records(all_rows)


@router.post("/import")
async def import_budgets(
    user: Annotated[CurrentUser, Depends(get_current_user)],
    db: Annotated[Client, Depends(get_db)],
    file: UploadFile,
):
    """Crée ou met à jour des budgets en lot depuis un .xlsx / .csv.
    Clé de rapprochement : catégorie + année + période (mois ou plage)."""
    _require_admin(user)
    raw = await file.read()
    name = (file.filename or "").lower()
    if name.endswith(".csv"):
        records = _parse_csv(raw)
    elif name.endswith((".xlsx", ".xlsm")):
        records = _parse_xlsx(raw)
    else:
        raise HTTPException(400, "Format non supporté — importez un fichier .xlsx ou .csv.")

    cats = db.from_("accounting_categories").select("id, name").execute().data or []
    cat_by_name = {_norm(c["name"]): c["id"] for c in cats}

    try:
        existing = db.from_("budgets").select(
            "id, category_id, year, month, start_date, end_date"
        ).execute().data or []
    except Exception as ex:
        raise _write_error(ex)
    idx = {(b["category_id"], b["year"], b["month"], b.get("start_date"), b.get("end_date")): b["id"]
           for b in existing}

    created = updated = 0
    errors: list[dict] = []
    for line, rec in enumerate(records, start=2):
        try:
            cat_raw = str(rec.get("category") or "").strip()
            if not cat_raw:
                continue
            cid = cat_by_name.get(_norm(cat_raw))
            if not cid:
                raise ValueError(f"catégorie introuvable : « {cat_raw} »")
            year_raw = rec.get("year")
            if year_raw in (None, ""):
                raise ValueError("année manquante")
            year = int(float(year_raw))
            m, s, e = _period_or_raise(
                _parse_month(rec.get("month")),
                _parse_date(rec.get("start_date")),
                _parse_date(rec.get("end_date")),
            )
            amount = _parse_amount(rec.get("amount"))
            comment = (str(rec.get("comment")).strip() if rec.get("comment") not in (None, "") else None)
            payload = {"category_id": cid, "year": year, "month": m,
                       "start_date": s, "end_date": e, "amount": amount, "comment": comment}
            key = (cid, year, m, s, e)
            if key in idx:
                db.from_("budgets").update(payload).eq("id", idx[key]).execute()
                updated += 1
            else:
                res = db.from_("budgets").insert({**payload, "created_by": user.id}).execute()
                idx[key] = res.data[0]["id"]
                created += 1
        except (ValueError, TypeError) as ex:
            errors.append({"row": line, "message": str(ex)})
        except Exception as ex:
            emsg = str(ex)
            if "does not exist" in emsg or "start_date" in emsg:
                raise _write_error(ex)
            errors.append({"row": line, "message": "erreur inattendue sur cette ligne"})

    log_audit(db, user.id, "budget.import", "budget", None,
              {"file": file.filename, "created": created, "updated": updated, "errors": len(errors)})
    return {"created": created, "updated": updated, "errors": errors,
            "total": len(records), "ok": created + updated}


@router.post("")
async def create_budget(
    body: BudgetCreate,
    user: Annotated[CurrentUser, Depends(get_current_user)],
    db: Annotated[Client, Depends(get_db)],
):
    _require_admin(user)
    data = body.model_dump()
    month, start_date, end_date = _clean_period(
        data.get("month"), data.get("start_date"), data.get("end_date"))
    payload = {
        "category_id": data["category_id"],
        "year": data["year"],
        "month": month,
        "start_date": start_date,
        "end_date": end_date,
        "amount": data.get("amount") or 0,
        "comment": data.get("comment"),
        "created_by": user.id,
    }
    try:
        res = db.from_("budgets").insert(payload).execute()
    except Exception as ex:
        raise _write_error(ex)
    new_budget = res.data[0]
    log_audit(db, user.id, "budget.create", "budget", new_budget["id"],
              {**payload, "created_by": None, "reference": new_budget.get("reference")})
    return _shape(new_budget)


@router.patch("/{budget_id}")
async def update_budget(
    budget_id: str,
    body: BudgetUpdate,
    user: Annotated[CurrentUser, Depends(get_current_user)],
    db: Annotated[Client, Depends(get_db)],
):
    _require_admin(user)
    raw = body.model_dump(exclude_unset=True)
    if not raw:
        raise HTTPException(400, "No fields to update")

    current = db.from_("budgets").select(
        "month, start_date, end_date").eq("id", budget_id).execute().data
    if not current:
        raise HTTPException(404, "Not found")
    cur = current[0]

    # `month` / `start_date` / `end_date` / `comment` peuvent être remis à NULL :
    # on les garde même quand la valeur envoyée est None (bascule de période).
    nullable = {"month", "start_date", "end_date", "comment"}
    updates = {k: v for k, v in raw.items() if k in nullable or v is not None}

    eff_month = updates["month"] if "month" in updates else cur.get("month")
    eff_start = updates["start_date"] if "start_date" in updates else cur.get("start_date")
    eff_end = updates["end_date"] if "end_date" in updates else cur.get("end_date")
    updates["month"], updates["start_date"], updates["end_date"] = _clean_period(
        eff_month, eff_start, eff_end)

    try:
        res = db.from_("budgets").update(updates).eq("id", budget_id).execute()
    except Exception as ex:
        raise _write_error(ex)
    if not res.data:
        raise HTTPException(404, "Not found")
    log_audit(db, user.id, "budget.update", "budget", budget_id, updates)
    return _shape(res.data[0])


@router.delete("/{budget_id}")
async def delete_budget(
    budget_id: str,
    user: Annotated[CurrentUser, Depends(get_current_user)],
    db: Annotated[Client, Depends(get_db)],
):
    _require_admin(user)
    existing = db.from_("budgets").select("id").eq("id", budget_id).execute().data
    if not existing:
        raise HTTPException(404, "Not found")
    db.from_("budgets").delete().eq("id", budget_id).execute()
    log_audit(db, user.id, "budget.delete", "budget", budget_id)
    return {"ok": True}
