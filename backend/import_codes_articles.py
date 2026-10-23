"""Import des codes articles (classeur « Codifications ») dans accounting_category_articles.

Chaque code article (ex. CM00012) est rattaché à la catégorie dont le code est son préfixe lettres
(CM → « Consommable Médical »). Une catégorie absente de la base est créée avec le nom du classeur.
Idempotent : un article est retrouvé par son code (mis à jour s'il change) ; rien n'est supprimé.
Essai à blanc par défaut ; --apply pour écrire.

    venv\Scripts\python.exe import_codes_articles.py "..\Codifications (1) (1).xlsx" [--apply]
"""
import argparse, re, sys
sys.stdout.reconfigure(encoding="utf-8")
from openpyxl import load_workbook

ap = argparse.ArgumentParser(); ap.add_argument("xlsx"); ap.add_argument("--apply", action="store_true"); args = ap.parse_args()
rows = list(load_workbook(args.xlsx, data_only=True)["Codes article"].iter_rows(min_row=2, values_only=True))
file_cats = {}
for r in rows:                                   # bloc catégories (colonnes A-C) — « Dépenses » prioritaire sur « Encaissements »
    if r[1] and (r[1] not in file_cats or r[0] == "Dépenses"):
        file_cats[str(r[1]).strip()] = str(r[2]).strip()
arts = []
for r in rows:
    if not r[5]: continue
    code = str(r[5]).strip()
    m = re.match(r"^([A-Z]+)\d+$", code)
    arts.append({"code": code, "prefix": m.group(1) if m else None, "article": re.sub(r"\s+", " ", str(r[6]).strip()),
                 "carac": (re.sub(r"\s+", " ", str(r[7]).strip()) if r[7] else None), "budget": r[8]})

from deps import get_db
db = get_db()
cats = {c["code"]: c for c in db.from_("accounting_categories").select("id,code,name").execute().data if c.get("code")}
existing = {a["code_article"]: a for a in db.from_("accounting_category_articles").select("*").execute().data if a.get("code_article")}
admin = (db.from_("user_roles").select("user_id").eq("role", "admin").execute().data or [{}])[0].get("user_id")
print(("== APPLICATION ==" if args.apply else "== ESSAI À BLANC ==") + f"  {len(arts)} articles dans le classeur")

new_cats, no_cat, created, updated, same = [], [], 0, 0, 0
for p in sorted({a["prefix"] for a in arts if a["prefix"]}):
    if p not in cats:
        name = file_cats.get(p)
        if not name: no_cat += [a["code"] for a in arts if a["prefix"] == p]; print(f"! préfixe {p} : aucune catégorie dans la base ni dans le classeur — {sum(1 for a in arts if a['prefix']==p)} articles ignorés"); continue
        new_cats.append((p, name)); print(f"+ catégorie {p} « {name} »")
        if args.apply: cats[p] = db.from_("accounting_categories").insert({"code": p, "name": name, "created_by": admin}).execute().data[0]
for a in arts:
    cat = cats.get(a["prefix"])
    comment = f"Budget indicatif : {a['budget']:g} MAD" if isinstance(a["budget"], (int, float)) else None
    row = {"code_article": a["code"], "article": a["article"], "caracteristiques": a["carac"], "commentaire": comment}
    if cat is None and not args.apply and a["prefix"] in dict(new_cats):
        cat = {"id": None}
    if cat is None: continue
    cur = existing.get(a["code"])
    if cur is None:
        created += 1
        if args.apply: db.from_("accounting_category_articles").insert({**row, "category_id": cat["id"], "created_by": admin}).execute()
    else:
        diff = {k: v for k, v in {**row, "category_id": cat["id"]}.items() if v is not None and cur.get(k) != v}
        if diff:
            updated += 1
            if "article" in diff: print(f"~ {a['code']} : « {cur['article']} » → « {a['article']} »")
            if args.apply: db.from_("accounting_category_articles").update(diff).eq("id", cur["id"]).execute()
        else: same += 1
print(f"\nRésumé : {created} créés, {updated} mis à jour, {same} inchangés, {len(no_cat)} ignorés, {len(new_cats)} catégories créées")
if not args.apply: print("Essai à blanc terminé — relancer avec --apply pour écrire.")
